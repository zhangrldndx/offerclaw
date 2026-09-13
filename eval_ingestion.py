# -*- coding: utf-8 -*-
"""eval_ingestion.py — 入库质量 Golden Set 基线(入库指导文档 §19 · 第 1 层采纳)。

动机:检索侧已有两把尺子,入库侧(PDF/Word → 文本)此前零评测。本脚本用
**确定性生成**的 5 类文档给当前 P0 抽取(文字层,零 OCR)建立基线;未来第 2 层
(Docling 结构化解析)按同口径 A/B——"测正才采纳"的前置条件。

Golden Set(reportlab/python-docx 代码生成,不提交二进制、完全可复现):
  1. 文字型 PDF(中文 CID 字体 + 英文各半——中文抽取质量本身就是待测项)
  2. 表格密集 PDF(platypus Table,已知单元格值)
  3. "扫描版" PDF(整页无文字层)          → 期望:明确拒收
  4. 含表格简历 DOCX(段落 + 技能表)
  5. 损坏文件(非 PDF 字节)               → 期望:可读报错拒收

指标(全部代码计算,零 LLM):
  parse_success / refusal_correct / coverage(已知句召回) /
  table_cell_recall(已知单元格召回) / docx_table_selfcontain(行级自含表头)

输出:docs/rag_eval/ingestion/baseline_p0.json(基线快照,供 Docling 对比)。
用法:.venv/bin/python eval_ingestion.py
"""
from __future__ import annotations

import io
import json
import os
import sys

BASE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, BASE)

OUT_DIR = os.path.join(BASE, "docs", "rag_eval", "ingestion")

# ---------------- 已知内容(金标准由代码保证) ----------------

CN_SENTENCES = [
    "OfferClaw 是一个长期运行的求职智能体系统。",
    "检索增强生成需要先建立可信的评测基线。",
    "知识库入库质量决定了检索层的上限。",
    "表格与图片是文档解析最常见的失真点。",
]
EN_SENTENCES = [
    "Retrieval quality depends on ingestion quality.",
    "Golden sets make parser upgrades measurable.",
    "Tables are where the most valuable data lives.",
    "Baselines must be recorded before optimization.",
]
TABLE_HEADERS = ["Skill", "Level", "Project"]
TABLE_ROWS = [
    ["LangGraph", "Familiar", "OfferClaw"],
    ["FastAPI", "Familiar", "LocalFlow"],
    ["ChromaDB", "Used", "OfferClaw"],
    ["Pytest", "Daily", "Both"],
]
DOCX_PARAS = [
    "张某,求职方向为大模型应用工程师。",
    "教育经历:某大学电子信息专业硕士在读。",
    "项目经历:主导构建了本地优先的求职智能体系统。",
]
DOCX_TABLE = [["技能", "熟练度"], ["LangGraph", "熟悉"], ["向量检索", "掌握"]]


# ---------------- Golden Set 生成(确定性) ----------------

def _cid_font():
    """reportlab 内置 Adobe CID 中文字体(无需外部字体文件)。"""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    return "STSong-Light"


def build_text_pdf() -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    font = _cid_font()
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    y = 780
    for s in CN_SENTENCES:
        c.setFont(font, 12)
        c.drawString(60, y, s)
        y -= 28
    for s in EN_SENTENCES:
        c.setFont("Helvetica", 12)
        c.drawString(60, y, s)
        y -= 28
    c.showPage()
    c.setFont(font, 12)
    c.drawString(60, 780, "第二页:跨页内容也应被完整抽取。")
    c.showPage()
    c.save()
    return buf.getvalue()


def build_table_pdf() -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Table
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4)
    doc.build([Table([TABLE_HEADERS] + TABLE_ROWS)])
    return buf.getvalue()


def build_scanned_pdf() -> bytes:
    """整页只有图形无文字层 = 扫描版等价物。"""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    c.rect(50, 400, 500, 300, fill=1)
    c.showPage()
    c.save()
    return buf.getvalue()


def build_resume_docx() -> bytes:
    from docx import Document
    buf = io.BytesIO()
    d = Document()
    for p in DOCX_PARAS:
        d.add_paragraph(p)
    t = d.add_table(rows=len(DOCX_TABLE), cols=2)
    for r, row in enumerate(DOCX_TABLE):
        for c_, v in enumerate(row):
            t.rows[r].cells[c_].text = v
    d.save(buf)
    return buf.getvalue()


def build_corrupt() -> bytes:
    return b"this is definitely not a valid pdf file"


# ---------------- 评测 ----------------

def _coverage(text: str, targets: list[str]) -> float:
    """已知句召回:空白归一后子串命中比例(CID 抽取可能引入空格差异)。"""
    import re
    norm = re.sub(r"\s+", "", text or "")
    hit = sum(1 for t in targets if re.sub(r"\s+", "", t) in norm)
    return round(hit / len(targets), 3) if targets else 1.0


def run_ingestion_eval() -> dict:
    from knowledge_crawler import extract_text_for_kb

    results: dict = {"parser": "P0 文字层(pypdf/python-docx,零 OCR 零 LLM)", "cases": {}}

    # 1. 文字型 PDF
    try:
        text = extract_text_for_kb("text.pdf", build_text_pdf())
        results["cases"]["text_pdf"] = {
            "parsed": True,
            "coverage_cn": _coverage(text, CN_SENTENCES + ["第二页:跨页内容也应被完整抽取。"]),
            "coverage_en": _coverage(text, EN_SENTENCES),
            "chars": len(text),
        }
    except ValueError as e:
        results["cases"]["text_pdf"] = {"parsed": False, "error": str(e)[:120]}

    # 2. 表格密集 PDF(P0 已知弱项:结构会碎)。诚实口径:
    #    cell_recall 只测"字符存在"(结构碎了也能满分),row_order 才测结构——
    #    同一行的值必须按序、且在近窗口内出现。这是留给第 2 层(Docling)对比的真靶子。
    try:
        text = extract_text_for_kb("table.pdf", build_table_pdf())
        cells = [v for row in TABLE_ROWS for v in row] + TABLE_HEADERS
        import re as _re
        norm = _re.sub(r"\s+", "", text or "")
        ordered = 0
        for row in TABLE_ROWS:
            idxs = [norm.find(v) for v in row]
            in_order = all(i >= 0 for i in idxs) and idxs == sorted(idxs)
            near = in_order and (idxs[-1] - idxs[0]) <= sum(len(v) for v in row) + 20
            ordered += 1 if (in_order and near) else 0
        results["cases"]["table_pdf"] = {
            "parsed": True,
            "table_cell_recall": _coverage(text, cells),
            "table_row_order": round(ordered / len(TABLE_ROWS), 3),
            "chars": len(text),
        }
    except ValueError as e:
        results["cases"]["table_pdf"] = {"parsed": False, "error": str(e)[:120]}

    # 3. 扫描版 → 期望拒收
    try:
        extract_text_for_kb("scan.pdf", build_scanned_pdf())
        results["cases"]["scanned_pdf"] = {"refused": False}
    except ValueError as e:
        results["cases"]["scanned_pdf"] = {"refused": True, "message_readable": "扫描" in str(e)}

    # 4. 含表格简历 DOCX
    try:
        text = extract_text_for_kb("resume.docx", build_resume_docx())
        results["cases"]["resume_docx"] = {
            "parsed": True,
            "coverage": _coverage(text, DOCX_PARAS),
            "table_cell_recall": _coverage(text, [v for r in DOCX_TABLE[1:] for v in r]),
            "table_row_selfcontain": ("技能: LangGraph；熟练度: 熟悉。" in text),
            "chars": len(text),
        }
    except ValueError as e:
        results["cases"]["resume_docx"] = {"parsed": False, "error": str(e)[:120]}

    # 5. 损坏文件 → 期望可读报错
    try:
        extract_text_for_kb("broken.pdf", build_corrupt())
        results["cases"]["corrupt"] = {"refused": False}
    except ValueError as e:
        results["cases"]["corrupt"] = {"refused": True, "message_readable": "解析失败" in str(e)}

    # 汇总
    c = results["cases"]
    results["summary"] = {
        "parse_success": sum(1 for k in ("text_pdf", "table_pdf", "resume_docx")
                             if c.get(k, {}).get("parsed")),
        "parse_expected": 3,
        "refusal_correct": sum(1 for k in ("scanned_pdf", "corrupt")
                               if c.get(k, {}).get("refused")),
        "refusal_expected": 2,
    }
    return results


def main() -> int:
    results = run_ingestion_eval()
    os.makedirs(OUT_DIR, exist_ok=True)
    out_path = os.path.join(OUT_DIR, "baseline_p0.json")
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(results, f, ensure_ascii=False, indent=2)
    print(json.dumps(results, ensure_ascii=False, indent=2))
    print(f"\n[baseline written] {os.path.relpath(out_path, BASE)}")
    s = results["summary"]
    ok = (s["parse_success"] == s["parse_expected"]
          and s["refusal_correct"] == s["refusal_expected"])
    return 0 if ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
