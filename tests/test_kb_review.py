# -*- coding: utf-8 -*-
"""知识库审核轨道测试:路径守卫 / 卡内预览 / 文件夹直投未入库扫描 / 疑似已在库标记 / API 防护。

隔离纪律:KB_DIR monkeypatch 到 tmp,不碰真实知识库;入库子进程与 chroma 全部打桩。
"""
from __future__ import annotations

import os
import sys

import pytest

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BASE)

import knowledge_crawler as kc  # noqa: E402


@pytest.fixture()
def kb_tmp(tmp_path, monkeypatch):
    """构造迷你知识库:正式两篇 + _pending 两篇(一篇与正式同标题)。"""
    kb = tmp_path / "knowledge_base"
    (kb / "learning_resources").mkdir(parents=True)
    (kb / "career_paths").mkdir()
    (kb / "_pending" / "web").mkdir(parents=True)
    (kb / "learning_resources" / "a_deep.md").write_text(
        '---\ntitle: "深度学习一"\n---\n# 正文A\n内容A', encoding="utf-8")
    (kb / "career_paths" / "b_path.md").write_text(
        '---\ntitle: "职业路线B"\n---\n# 正文B\n内容B', encoding="utf-8")
    (kb / "_pending" / "web" / "c_dup.md").write_text(
        '---\ntitle: "深度学习一"\n---\n重复暂存副本', encoding="utf-8")
    (kb / "_pending" / "web" / "d_new.md").write_text(
        '---\ntitle: "全新候选D"\n---\n新内容', encoding="utf-8")
    monkeypatch.setattr(kc, "KB_DIR", str(kb))
    return kb


# ---------- 路径守卫(预览/入库的唯一入口) ----------

def test_safe_kb_file_accepts_and_normalizes(kb_tmp):
    p1 = kc.safe_kb_file("learning_resources/a_deep.md")
    p2 = kc.safe_kb_file("knowledge_base/learning_resources/a_deep.md")  # 带前缀同样接受
    assert p1 == p2 and p1.endswith("a_deep.md")
    assert kc.safe_kb_file("_pending/web/c_dup.md").endswith("c_dup.md")


@pytest.mark.parametrize("bad", [
    "../secrets.txt",                       # 越界
    "learning_resources/../../x.md",        # 变形越界
    "learning_resources/nope.md",           # 不存在
    "",                                     # 空
])
def test_safe_kb_file_rejects(kb_tmp, bad):
    with pytest.raises(ValueError):
        kc.safe_kb_file(bad)


def test_safe_kb_file_rejects_non_text(kb_tmp):
    (kb_tmp / "learning_resources" / "img.png").write_bytes(b"\x89PNG")
    with pytest.raises(ValueError):
        kc.safe_kb_file("learning_resources/img.png")


# ---------- 卡内预览 ----------

def test_read_kb_preview_and_truncation(kb_tmp):
    d = kc.read_kb_preview("learning_resources/a_deep.md")
    assert d["title"] == "深度学习一" and "正文A" in d["content"] and d["truncated"] is False
    t = kc.read_kb_preview("learning_resources/a_deep.md", max_chars=10)
    assert t["truncated"] is True and len(t["content"]) == 10


# ---------- 文件夹直投:未入库扫描 ----------

def test_list_unindexed_diffs_disk_vs_index(kb_tmp):
    items = kc.list_unindexed({"a_deep.md"})          # a 已入索引
    rels = {it["rel"] for it in items}
    assert rels == {"career_paths/b_path.md"}         # 只剩 b;_pending 不参与
    assert items[0]["subdir"] == "career_paths"
    assert kc.list_unindexed({"a_deep.md", "b_path.md"}) == []


# ---------- 疑似已在库标记 ----------

def test_pending_archive_dir_hidden_from_candidates(kb_tmp):
    """_pending/_archived/ 归档区不进待审列表(清理 83 条遗留副本的通道),文件本体保留。"""
    arch = kb_tmp / "_pending" / "_archived"
    arch.mkdir()
    (arch / "old_dup.md").write_text(
        '---\ntitle: "旧副本"\nsource_url: "https://x"\n---\n内容', encoding="utf-8")
    (kb_tmp / "_pending" / "web" / "live.md").write_text(
        '---\ntitle: "活跃候选"\nsource_url: "https://y"\n---\n内容', encoding="utf-8")
    titles = [it["title"] for it in kc.cmd_list_pending()["items"]]
    assert "活跃候选" in titles and "旧副本" not in titles
    assert (arch / "old_dup.md").exists()


def test_flag_existing_in_kb(kb_tmp):
    items = [{"title": "深度学习一"}, {"title": "全新候选D"}, {"title": ""}]
    kc.flag_existing_in_kb(items)
    assert items[0]["existing_in_kb"] is True         # 与正式库同标题 → 疑似重复暂存
    assert items[1]["existing_in_kb"] is False
    assert items[2]["existing_in_kb"] is False


# ---------- API 层防护 ----------

@pytest.fixture()
def client():
    from fastapi.testclient import TestClient
    import rag_api
    return TestClient(rag_api.app)


def test_api_preview_guards_traversal(client, kb_tmp):
    assert client.get("/api/kb/preview", params={"rel": "../metrics.json"}).status_code == 400
    r = client.get("/api/kb/preview", params={"rel": "learning_resources/a_deep.md"})
    assert r.status_code == 200 and r.json()["title"] == "深度学习一"


def test_api_ingest_path_rejects_pending_and_traversal(client, kb_tmp, monkeypatch):
    assert client.post("/api/kb/ingest_path", json={"rel": "../x.md"}).status_code == 400
    r = client.post("/api/kb/ingest_path", json={"rel": "_pending/web/c_dup.md"})
    assert r.status_code == 400 and "promote" in r.json()["detail"]


def test_api_ingest_path_happy_with_stubbed_subprocess(client, kb_tmp, monkeypatch):
    import subprocess as sp
    import rag_api

    class _Ok:
        returncode = 0
        stdout = "ok"
        stderr = ""

    monkeypatch.setattr(sp, "run", lambda *a, **k: _Ok())
    counts = iter([100, 103])
    monkeypatch.setattr(rag_api, "_kb_count", lambda: next(counts))
    monkeypatch.setattr(rag_api, "_kb_clear_cache", lambda: None)
    r = client.post("/api/kb/ingest_path", json={"rel": "career_paths/b_path.md"})
    assert r.status_code == 200
    j = r.json()
    assert j["chunks_added"] == 3 and j["file"].endswith("b_path.md")


def test_api_unindexed_shape(client, kb_tmp, monkeypatch):
    import rag_api
    monkeypatch.setattr(rag_api, "_indexed_sources", lambda: {"a_deep.md"})
    r = client.get("/api/kb/unindexed")
    assert r.status_code == 200
    j = r.json()
    assert j["count"] == 1 and j["items"][0]["file"] == "b_path.md"


# ---------- PDF / Word 上传抽取(途径B格式扩展,2026-08-08) ----------

def test_extract_docx_paragraphs_and_tables(tmp_path):
    from docx import Document
    p = tmp_path / "t.docx"
    doc = Document()
    doc.add_paragraph("这是一段足够长的正文内容,用于验证抽取管线是否工作。")
    doc.add_paragraph("第二段:RAG 知识库测试样例文字若干。")
    t = doc.add_table(rows=1, cols=2)          # 单行表:无表头语义,保持竖线拼接
    t.rows[0].cells[0].text = "表格甲"
    t.rows[0].cells[1].text = "表格乙"
    t2 = doc.add_table(rows=3, cols=2)          # 多行表:句子化,行级自含表头(§9.2)
    for r, (a, b) in enumerate([("技能", "熟练度"), ("LangGraph", "熟悉"), ("FastAPI", "掌握")]):
        t2.rows[r].cells[0].text = a
        t2.rows[r].cells[1].text = b
    doc.save(str(p))
    text = kc.extract_text_for_kb("t.docx", p.read_bytes())
    assert "第二段" in text and "表格甲 | 表格乙" in text
    assert "表头：技能、熟练度" in text
    assert "技能: LangGraph；熟练度: 熟悉。" in text     # 每行自带表头 → 分块边界安全
    assert "技能: FastAPI；熟练度: 掌握。" in text


def test_duplicate_content_detection(kb_tmp):
    """文件级内容去重(§5.1):同内容(空白重排/带 frontmatter)命中,改写不命中。"""
    body = "内容A"                                        # kb_tmp 里 a_deep.md 的正文
    hit = kc.find_duplicate_content(f"# 正文A\n{body}")
    assert hit and hit["where"] == "formal" and hit["title"] == "深度学习一"
    # 空白重排仍命中;frontmatter 剥离后比对
    assert kc.find_duplicate_content("---\ntitle: x\n---\n# 正文A\n\n\n  内容A  ")
    assert kc.find_duplicate_content("完全不同的新内容,不应命中任何已有文件。") is None


def test_score_and_save_skips_duplicate(kb_tmp, monkeypatch):
    """重复上传:打分函数不被调用(省 LLM),返回 duplicate 状态与已有文件指引。"""
    monkeypatch.setattr(kc, "PENDING_WEB_DIR", str(kb_tmp / "_pending" / "web"))
    called = {"n": 0}

    def _boom(text, title=""):
        called["n"] += 1
        raise AssertionError("重复内容不应触发打分")

    monkeypatch.setattr(kc, "score_content", _boom)
    out = kc._score_and_save("# 正文A\n内容A", url="(本地上传:dup.md)", origin="本地上传",
                             force_keep=True)
    assert out["status"] == "duplicate" and out["existing"]["where"] == "formal"
    assert called["n"] == 0


def test_extract_blank_pdf_rejected():
    import io
    from pypdf import PdfWriter
    buf = io.BytesIO()
    w = PdfWriter()
    w.add_blank_page(width=200, height=200)
    w.write(buf)
    with pytest.raises(ValueError):        # 无文字层(等价扫描版) → 明确拒绝,不静默入库空壳
        kc.extract_text_for_kb("scan.pdf", buf.getvalue())


def test_extract_unsupported_ext_rejected():
    with pytest.raises(ValueError):
        kc.extract_text_for_kb("x.doc", b"xx")   # 旧版 .doc 不支持


def test_api_upload_docx_lands_in_pending(client, kb_tmp, monkeypatch, tmp_path):
    import base64
    from docx import Document
    monkeypatch.setattr(kc, "PENDING_WEB_DIR", str(kb_tmp / "_pending" / "web"))
    monkeypatch.setattr(kc, "score_content", lambda text, title="": {
        "grade": "B", "relevance": 7, "density": 7, "recency_ok": True,
        "reason": "stub", "suggested_subdir": "learning_resources",
        "suggested_title": "上传测试"})
    p = tmp_path / "u.docx"
    d = Document()
    # _score_and_save 有"正文 <80 字符即拒收"的质量门,语料须够长
    d.add_paragraph("上传的 Word 文档正文:本段用于验证 PDF/Word 抽取管线端到端工作,"
                    "包含足够的字符数以通过候选落盘前的最小长度质量门,并附带若干技术关键词:"
                    "RAG、向量检索、增量索引、frontmatter 元数据与人工审核流程。")
    d.save(str(p))
    b64 = base64.b64encode(p.read_bytes()).decode()
    r = client.post("/api/kb/add_file", json={"name": "u.docx", "content_base64": b64})
    assert r.status_code == 200 and r.json().get("status") != "rejected"
    files = list((kb_tmp / "_pending" / "web").glob("*.md"))
    bodies = [f.read_text(encoding="utf-8") for f in files]
    assert any("上传的 Word 文档正文" in b and "source_url" in b for b in bodies), \
        "抽取文本应作为带 source_url 元数据的候选落盘 _pending/web"


def test_api_upload_scan_pdf_400(client, kb_tmp):
    import base64
    import io
    from pypdf import PdfWriter
    buf = io.BytesIO()
    w = PdfWriter()
    w.add_blank_page(width=100, height=100)
    w.write(buf)
    r = client.post("/api/kb/add_file",
                    json={"name": "s.pdf", "parser": "text",   # 显式文字层(自动路由下 PDF 默认走结构化)
                          "content_base64": base64.b64encode(buf.getvalue()).decode()})
    assert r.status_code == 400 and "扫描" in r.json()["detail"]


# ---------- Docling opt-in 分发 + 论文域(2026-08-09) ----------

def test_papers_subdir_registered():
    assert "papers" in kc.VALID_SUBDIRS and kc.SUBDIR_SOURCE_TYPE["papers"] == "paper"
    from rag_ingest import _infer_source_type
    assert _infer_source_type("knowledge_base/papers/x.md") == "paper"


def test_api_add_file_parser_dispatch(client, kb_tmp, monkeypatch):
    """parser=docling 走结构化引擎(打桩);非法 parser 400;默认路径不受影响。"""
    import base64
    monkeypatch.setattr(kc, "PENDING_WEB_DIR", str(kb_tmp / "_pending" / "web"))
    monkeypatch.setattr(kc, "score_content", lambda text, title="": {
        "grade": "B", "relevance": 7, "density": 7, "recency_ok": True,
        "reason": "stub", "suggested_subdir": "papers", "suggested_title": "论文测试"})
    monkeypatch.setattr(kc, "extract_text_structured",
                        lambda name, raw: "## 结构化标题\n\n| 表头 | 值 |\n|---|---|\n| a | b |\n\n"
                                          "正文内容足够长以通过候选落盘前的最小长度质量门,附加填充文字确保超过八十个字符的下限要求。")
    b64 = base64.b64encode(b"%PDF-fake").decode()
    r = client.post("/api/kb/add_file",
                    json={"name": "paper.pdf", "content_base64": b64, "parser": "docling"})
    assert r.status_code == 200 and r.json().get("status") != "rejected"
    bodies = [f.read_text(encoding="utf-8")
              for f in (kb_tmp / "_pending" / "web").glob("*.md")]
    assert any("## 结构化标题" in b for b in bodies)      # 用的是结构化引擎输出
    assert client.post("/api/kb/add_file",
                       json={"name": "x.pdf", "content_base64": b64,
                             "parser": "magic"}).status_code == 400


def test_structured_engine_rejects_non_pdf_docx(kb_tmp):
    with pytest.raises(ValueError):
        kc.extract_text_structured("x.md", b"text")


def test_api_pdf_auto_routes_to_structured(client, kb_tmp, monkeypatch):
    """自动路由(2026-08-09):PDF 不带 parser 参数默认走结构化引擎,docx 默认原生。"""
    import base64
    monkeypatch.setattr(kc, "PENDING_WEB_DIR", str(kb_tmp / "_pending" / "web"))
    monkeypatch.setattr(kc, "score_content", lambda text, title="": {
        "grade": "B", "relevance": 7, "density": 7, "recency_ok": True,
        "reason": "stub", "suggested_subdir": "papers", "suggested_title": "自动路由"})
    monkeypatch.setattr(kc, "extract_text_structured",
                        lambda name, raw: "## 自动结构化输出\n\n正文足够长以通过质量门,"
                                          "为通过八十字符质量门补足的填充语料:检索增强生成、向量数据库、混合检索、精排、拒答门控、评测基线、增量索引、内容去重、结构化解析。")
    b64 = base64.b64encode(b"%PDF-fake").decode()
    r = client.post("/api/kb/add_file", json={"name": "auto.pdf", "content_base64": b64})
    assert r.status_code == 200 and r.json().get("status") != "rejected"
    bodies = [f.read_text(encoding="utf-8")
              for f in (kb_tmp / "_pending" / "web").glob("*.md")]
    assert any("## 自动结构化输出" in b for b in bodies)


def test_api_pdf_falls_back_to_text_when_structured_fails(client, kb_tmp, monkeypatch):
    """自动路由的回退:结构化失败 → 文字层兜底(显式指定 docling 时才 fail-visible)。"""
    import base64
    monkeypatch.setattr(kc, "PENDING_WEB_DIR", str(kb_tmp / "_pending" / "web"))
    monkeypatch.setattr(kc, "score_content", lambda text, title="": {
        "grade": "C", "relevance": 5, "density": 5, "recency_ok": True,
        "reason": "stub", "suggested_subdir": "learning_resources", "suggested_title": "回退"})
    def _boom(name, raw):
        raise ValueError("模拟结构化不可用")
    monkeypatch.setattr(kc, "extract_text_structured", _boom)
    monkeypatch.setattr(kc, "extract_text_for_kb",
                        lambda name, raw: "文字层兜底输出:内容足够长以通过候选落盘前的最小"
                                          "为通过八十字符质量门补足的填充语料:检索增强生成、向量数据库、混合检索、精排、拒答门控、评测基线、增量索引、内容去重、结构化解析。")
    b64 = base64.b64encode(b"%PDF-fake").decode()
    r = client.post("/api/kb/add_file", json={"name": "fb.pdf", "content_base64": b64})
    assert r.status_code == 200
    # 显式 docling 则不回退,原样报错
    r2 = client.post("/api/kb/add_file",
                     json={"name": "fb2.pdf", "content_base64": b64, "parser": "docling"})
    assert r2.status_code == 400 and "模拟结构化不可用" in r2.json()["detail"]
